"""Daily report + XLSX export."""

import io
import re
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from lib.auth import current_user
from lib.catalog import stage_labels
from lib.dates import today_iso
from lib.permissions import require
from lib.db import db
from models.schemas import (
    STAGE_LABELS,
    Activity,
    CurrencyTotal,
    DailyReport,
    DealerMonthRow,
    ExchangeRate,
    MonthlyReport,
    Project,
    StageCount,
)

router = APIRouter(prefix="/reports", tags=["reports"])


def _by_currency(projects: list[Project]) -> list[CurrencyTotal]:
    """Para birimi bazında kırılım — farklı kurlardaki tutarlar asla tek toplamda birleşmez."""
    groups: dict[str, list[Project]] = {}
    for p in projects:
        groups.setdefault(p.para_birimi or "—", []).append(p)
    rows = [
        CurrencyTotal(
            para_birimi=cur,
            proje_adet=len(group),
            satis=round(sum(p.muhasebe.transfer_dahil_toplam_satis for p in group), 2),
            tahsilat=round(sum(p.muhasebe.toplam_tahsilat for p in group), 2),
            bakiye=round(sum(p.muhasebe.kalan_bakiye for p in group), 2),
            net_kar=round(sum(p.muhasebe.net_kar for p in group), 2),
        )
        for cur, group in groups.items()
    ]
    return sorted(rows, key=lambda r: r.satis, reverse=True)


def _aware(doc: dict) -> dict:
    for key in ("created_at", "updated_at"):
        value = doc.get(key)
        if isinstance(value, datetime) and value.tzinfo is None:
            doc[key] = value.replace(tzinfo=timezone.utc)
    return doc


async def _collect(tarih: str) -> DailyReport:
    acts = [
        Activity(**_aware(a))
        for a in await db.activities.find({"gun": tarih}).sort("created_at", -1).to_list(500)
    ]
    yeni_docs = await db.projects.find({"proje_tarihi": tarih}).to_list(200)
    sevk_docs = await db.projects.find({"sevk_tarihi": tarih}).to_list(200)
    aktif = await db.projects.count_documents({"arsiv": False})
    yeni = [Project(**_aware(d)) for d in yeni_docs]
    sevk = [Project(**_aware(d)) for d in sevk_docs]
    return DailyReport(
        tarih=tarih,
        yeni_projeler=yeni,
        asama_degisimleri=[a for a in acts if a.tip == "asama"],
        sevk_edilenler=sevk,
        tum_hareketler=acts,
        gun_satis=round(sum(p.muhasebe.transfer_dahil_toplam_satis for p in yeni), 2),
        gun_tahsilat=round(sum(p.muhasebe.toplam_tahsilat for p in sevk), 2),
        aktif_proje=aktif,
        gun_satis_dagilimi=_by_currency(yeni),
        gun_tahsilat_dagilimi=_by_currency(sevk),
        genel_dagilim=_by_currency(
            [Project(**_aware(d)) for d in await db.projects.find({"arsiv": False}).to_list(1000)]
        ),
    )


@router.get("/dealer-statement")
async def dealer_statement(
    anahtar: str = Query(...), user: dict = Depends(require("rapor:disaari"))
):
    """Bir bayinin tüm projeleri + tahsilat dökümü — tek Excel dosyası."""
    firma, _, ulke = anahtar.partition("|")
    query: dict = {"firma": firma}
    if ulke:
        query["ulke"] = ulke
    docs = await db.projects.find(query).sort("proje_tarihi", -1).to_list(500)
    if not docs:
        raise HTTPException(status_code=404, detail="Bayi bulunamadı")
    projects = [Project(**_aware(d)) for d in docs]
    labels = await stage_labels()
    ids = [p.id for p in projects]
    code_by_id = {p.id: p.proje_kodu for p in projects}
    items = await db.project_items.find({"proje_id": {"$in": ids}}).to_list(2000)
    crates = await db.project_crates.find({"proje_id": {"$in": ids}}).to_list(2000)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F172A")
    title_font = Font(bold=True, size=13)

    def sheet(title: str, headers: list[str], rows: list[list], baslik: str | None = None):
        ws = wb.create_sheet(title)
        if baslik:
            ws.append([baslik])
            ws["A1"].font = title_font
            ws.append([])
        ws.append(headers)
        for cell in ws[ws.max_row]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
        for idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(
                14, min(40, len(header) + 6)
            )
        return ws

    wb.remove(wb.active)

    ciro = round(sum(p.muhasebe.transfer_dahil_toplam_satis for p in projects), 2)
    tahsilat = round(sum(p.muhasebe.toplam_tahsilat for p in projects), 2)
    sheet(
        "Bayi Ozeti",
        [
            "Firma",
            "Ulke",
            "Musteriler",
            "Proje Adedi",
            "Aktif",
            "Arsiv",
            "Ciro",
            "Tahsilat",
            "Acik Bakiye",
            "Net Kar",
            "Kar %",
        ],
        [
            [
                firma,
                ulke,
                ", ".join(sorted({p.musteri for p in projects if p.musteri})),
                len(projects),
                len([p for p in projects if not p.arsiv]),
                len([p for p in projects if p.arsiv]),
                ciro,
                tahsilat,
                round(sum(p.muhasebe.kalan_bakiye for p in projects), 2),
                round(sum(p.muhasebe.net_kar for p in projects), 2),
                round(sum(p.muhasebe.net_kar for p in projects) / ciro * 100, 2) if ciro else 0,
            ]
        ],
        baslik=f"{firma} — {ulke} Bayi Ekstresi",
    )

    sheet(
        "Projeler",
        [
            "Proje Kodu",
            "Proje Adi",
            "Musteri",
            "Tarih",
            "Asama",
            "Tedarikci",
            "Para Birimi",
            "Satis",
            "Iskonto",
            "Transfer",
            "Toplam Satis",
            "Alis",
            "Net Kar",
            "Kar %",
            "Tahsilat",
            "Kalan Bakiye",
            "Odeme Durumu",
            "Sevk Tarihi",
            "Arsiv",
        ],
        [
            [
                p.proje_kodu,
                p.proje_adi,
                p.musteri,
                p.proje_tarihi,
                labels.get(p.durum, p.durum),
                p.tedarikci,
                p.para_birimi,
                p.muhasebe.satis,
                p.muhasebe.iskonto_tutari,
                p.muhasebe.transfer_ucreti,
                p.muhasebe.transfer_dahil_toplam_satis,
                p.muhasebe.alis,
                p.muhasebe.net_kar,
                p.muhasebe.kar_yuzdesi,
                p.muhasebe.toplam_tahsilat,
                p.muhasebe.kalan_bakiye,
                p.muhasebe.odeme_durumu,
                p.sevk_tarihi or "",
                "Evet" if p.arsiv else "Hayir",
            ]
            for p in projects
        ],
    )

    tahsilat_rows: list[list] = []
    for p in projects:
        for i, o in enumerate(p.muhasebe.odemeler, start=1):
            if not o.tutar:
                continue
            tahsilat_rows.append(
                [p.proje_kodu, p.musteri, f"{i}. Odeme", o.tutar, o.tarih or "", p.para_birimi, o.not_]
            )
        tahsilat_rows.append(
            [
                p.proje_kodu,
                p.musteri,
                "KALAN BAKIYE",
                p.muhasebe.kalan_bakiye,
                "",
                p.para_birimi,
                p.muhasebe.odeme_durumu,
            ]
        )
    sheet(
        "Tahsilat Dokumu",
        ["Proje Kodu", "Musteri", "Kalem", "Tutar", "Tarih", "Para Birimi", "Not"],
        tahsilat_rows,
    )

    sheet(
        "Proje Kalemleri",
        [
            "Proje Kodu",
            "Urun",
            "Adet",
            "Genislik (mm)",
            "Acilim (mm)",
            "Yapi Rengi",
            "Panel Rengi",
            "Aydinlatma",
            "Cam Kombinasyonu",
            "Tedarikci",
            "Birim Fiyat",
            "Tutar",
        ],
        [
            [
                code_by_id.get(i.get("proje_id", ""), ""),
                i.get("urun", ""),
                i.get("adet", 0),
                i.get("genislik_mm", 0),
                i.get("acilim_mm", 0),
                i.get("yapi_rengi", ""),
                i.get("panel_rengi", ""),
                i.get("aydinlatma", ""),
                i.get("cam_kombinasyonu", ""),
                i.get("tedarikci", ""),
                i.get("birim_fiyat", 0),
                (i.get("adet") or 0) * (i.get("birim_fiyat") or 0),
            ]
            for i in items
        ],
    )

    sheet(
        "Sandik Listesi",
        [
            "Proje Kodu",
            "Sandik No",
            "Icerik",
            "Taban (cm)",
            "Uzunluk (cm)",
            "Yukseklik (cm)",
            "Adet",
            "Hacim (m3)",
            "Brut (kg)",
            "Tedarikci",
        ],
        [
            [
                code_by_id.get(c.get("proje_id", ""), ""),
                c.get("sandik_no", ""),
                c.get("icerik", ""),
                c.get("taban_cm", 0),
                c.get("uzunluk_cm", 0),
                c.get("yukseklik_cm", 0),
                c.get("adet", 0),
                c.get("hacim_cbm", 0),
                c.get("brut_kg", 0),
                c.get("tedarikci", ""),
            ]
            for c in crates
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = (
        firma.replace("ı", "i")
        .replace("İ", "I")
        .replace("ş", "s")
        .replace("Ş", "S")
        .replace("ğ", "g")
        .replace("Ğ", "G")
        .replace("ü", "u")
        .replace("Ü", "U")
        .replace("ö", "o")
        .replace("Ö", "O")
        .replace("ç", "c")
        .replace("Ç", "C")
        .encode("ascii", "ignore")
        .decode()
    )
    safe = "".join(ch for ch in safe if ch.isalnum() or ch in " -_").strip() or "bayi"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="bayi-ekstre-{safe}.xlsx"'},
    )


@router.get("/daily", response_model=DailyReport)
async def daily_report(
    tarih: Optional[str] = Query(default=None), user: dict = Depends(require("rapor:goruntule"))
):
    return await _collect(tarih or today_iso())


@router.get("/daily/export")
async def export_daily(
    tarih: Optional[str] = Query(default=None), user: dict = Depends(require("rapor:disaari"))
):
    gun = tarih or today_iso()
    report = await _collect(gun)
    projects = [
        Project(**_aware(d))
        for d in await db.projects.find().sort("proje_tarihi", -1).to_list(1000)
    ]

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F172A")

    def sheet(title: str, headers: list[str], rows: list[list]):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
        for idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(
                14, min(38, len(header) + 6)
            )
        return ws

    wb.remove(wb.active)

    sheet(
        "Gun Ozeti",
        ["Rapor Tarihi", "Yeni Proje", "Asama Degisimi", "Sevk Edilen", "Aktif Proje"],
        [
            [
                gun,
                len(report.yeni_projeler),
                len(report.asama_degisimleri),
                len(report.sevk_edilenler),
                report.aktif_proje,
            ]
        ],
    )

    # Para birimi kırılımı — farklı kurlar tek toplamda birleştirilmez.
    sheet(
        "Para Birimi Dagilimi",
        [
            "Kapsam",
            "Para Birimi",
            "Proje Adedi",
            "Toplam Satis",
            "Tahsilat",
            "Kalan Bakiye",
            "Net Kar",
        ],
        [
            [kapsam, r.para_birimi, r.proje_adet, r.satis, r.tahsilat, r.bakiye, r.net_kar]
            for kapsam, rows in (
                ("Gun Ici Acilan", report.gun_satis_dagilimi),
                ("Gun Ici Sevk Edilen", report.gun_tahsilat_dagilimi),
                ("Tum Aktif Projeler", report.genel_dagilim),
            )
            for r in rows
        ],
    )

    sheet(
        "Gun Hareketleri",
        ["Saat", "Proje Kodu", "Tip", "Aciklama", "Kullanici"],
        [
            [a.created_at.strftime("%H:%M"), a.proje_kodu, a.tip, a.mesaj, a.kullanici]
            for a in report.tum_hareketler
        ],
    )

    sheet(
        "Projeler",
        [
            "Proje Kodu",
            "Firma",
            "Musteri",
            "Ulke",
            "Proje Adi",
            "Tarih",
            "Asama",
            "Tedarikci",
            "Para Birimi",
            "Satis (PB)",
            "Alis (PB)",
            "Iskonto (PB)",
            "Transfer (PB)",
            "Toplam Satis (PB)",
            "Net Kar (PB)",
            "Kar %",
            "Odeme Durumu",
            "Tahsilat (PB)",
            "Kalan Bakiye (PB)",
            "Sevk Tarihi",
            "Arsiv",
        ],
        [
            [
                p.proje_kodu,
                p.firma,
                p.musteri,
                p.ulke,
                p.proje_adi,
                p.proje_tarihi,
                STAGE_LABELS.get(p.durum, p.durum),
                p.tedarikci,
                p.para_birimi,
                f"{p.muhasebe.satis:.2f} {p.para_birimi}".strip(),
                f"{p.muhasebe.alis:.2f} {p.para_birimi}".strip(),
                f"{p.muhasebe.iskonto_tutari:.2f} {p.para_birimi}".strip(),
                f"{p.muhasebe.transfer_ucreti:.2f} {p.para_birimi}".strip(),
                f"{p.muhasebe.transfer_dahil_toplam_satis:.2f} {p.para_birimi}".strip(),
                f"{p.muhasebe.net_kar:.2f} {p.para_birimi}".strip(),
                p.muhasebe.kar_yuzdesi,
                p.muhasebe.odeme_durumu,
                f"{p.muhasebe.toplam_tahsilat:.2f} {p.para_birimi}".strip(),
                f"{p.muhasebe.kalan_bakiye:.2f} {p.para_birimi}".strip(),
                p.sevk_tarihi or "",
                "Evet" if p.arsiv else "Hayir",
            ]
            for p in projects
        ],
    )

    items = await db.project_items.find().to_list(2000)
    code_by_id = {p.id: p.proje_kodu for p in projects}
    cur_by_id = {p.id: p.para_birimi for p in projects}
    sheet(
        "Proje Kalemleri",
        [
            "Proje Kodu",
            "Urun",
            "Adet",
            "Genislik (mm)",
            "Acilim (mm)",
            "Yapi Rengi",
            "Panel Rengi",
            "Aydinlatma",
            "Cam Kombinasyonu",
            "Tedarikci",
            "Para Birimi",
            "Birim Fiyat (PB)",
            "Satir Tutari (PB)",
        ],
        [
            [
                code_by_id.get(i.get("proje_id", ""), ""),
                i.get("urun", ""),
                i.get("adet", 0),
                i.get("genislik_mm", 0),
                i.get("acilim_mm", 0),
                i.get("yapi_rengi", ""),
                i.get("panel_rengi", ""),
                i.get("aydinlatma", ""),
                i.get("cam_kombinasyonu", ""),
                i.get("tedarikci", ""),
                cur_by_id.get(i.get("proje_id", ""), ""),
                f"{float(i.get('birim_fiyat') or 0):.2f} "
                f"{cur_by_id.get(i.get('proje_id', ''), '')}".strip(),
                f"{float(i.get('birim_fiyat') or 0) * float(i.get('adet') or 0):.2f} "
                f"{cur_by_id.get(i.get('proje_id', ''), '')}".strip(),
            ]
            for i in items
        ],
    )

    crates = await db.project_crates.find().to_list(2000)
    sheet(
        "Sandik Listesi",
        [
            "Proje Kodu",
            "Sandik No",
            "Icerik",
            "Taban (cm)",
            "Uzunluk (cm)",
            "Yukseklik (cm)",
            "Adet",
            "Hacim (m3)",
            "Brut (kg)",
            "Tedarikci",
        ],
        [
            [
                code_by_id.get(c.get("proje_id", ""), ""),
                c.get("sandik_no", ""),
                c.get("icerik", ""),
                c.get("taban_cm", 0),
                c.get("uzunluk_cm", 0),
                c.get("yukseklik_cm", 0),
                c.get("adet", 0),
                c.get("hacim_cbm", 0),
                c.get("brut_kg", 0),
                c.get("tedarikci", ""),
            ]
            for c in crates
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="pergola-rapor-{gun}.xlsx"'},
    )


# ---------- kur çevirisi + aylık rapor ----------
async def _rates() -> dict[str, float]:
    """TRY bazlı kurlar: 1 birim = kaç TRY. TRY her zaman 1.0."""
    rows = await db.rates.find().to_list(50)
    rates = {r["para_birimi"].upper(): float(r.get("kur") or 0) for r in rows}
    rates["TRY"] = 1.0
    return rates


def _to_try(amount: float, currency: str, rates: dict[str, float]) -> float:
    return amount * rates.get((currency or "TRY").upper(), 0.0)


@router.get("/monthly", response_model=MonthlyReport)
async def monthly_report(
    ay: Optional[str] = Query(default=None, description="YYYY-MM"),
    user: dict = Depends(require("rapor:goruntule")),
):
    donem = ay or today_iso()[:7]
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", donem):
        raise HTTPException(status_code=400, detail="Dönem YYYY-AA biçiminde olmalı")
    docs = await db.projects.find({"proje_tarihi": {"$regex": f"^{re.escape(donem)}"}}).to_list(2000)
    projects = [Project(**_aware(d)) for d in docs]

    rates = await _rates()
    rate_rows = [ExchangeRate(**r) for r in await db.rates.find().sort("para_birimi", 1).to_list(50)]
    kullanilan = {(p.para_birimi or "TRY").upper() for p in projects}
    eksik = sorted(c for c in kullanilan if not rates.get(c))

    labels = await stage_labels()
    asama = [
        StageCount(
            durum=d,
            label=labels.get(d, d),
            adet=len([p for p in projects if p.durum == d]),
            tutar=round(
                sum(
                    _to_try(p.muhasebe.transfer_dahil_toplam_satis, p.para_birimi, rates)
                    for p in projects
                    if p.durum == d
                ),
                2,
            ),
        )
        for d in labels
        if any(p.durum == d for p in projects)
    ]

    bayiler: dict[str, list[Project]] = {}
    for p in projects:
        bayiler.setdefault(f"{p.firma or 'Belirtilmemiş'}|{p.ulke or '—'}", []).append(p)
    bayi_ozeti = sorted(
        [
            DealerMonthRow(
                firma=key.split("|", 1)[0],
                ulke=key.split("|", 1)[1],
                proje_adet=len(group),
                satis_try=round(
                    sum(
                        _to_try(p.muhasebe.transfer_dahil_toplam_satis, p.para_birimi, rates)
                        for p in group
                    ),
                    2,
                ),
                tahsilat_try=round(
                    sum(_to_try(p.muhasebe.toplam_tahsilat, p.para_birimi, rates) for p in group), 2
                ),
            )
            for key, group in bayiler.items()
        ],
        key=lambda r: r.satis_try,
        reverse=True,
    )

    return MonthlyReport(
        ay=donem,
        proje_adet=len(projects),
        kur_dagilimi=_by_currency(projects),
        try_satis=round(
            sum(
                _to_try(p.muhasebe.transfer_dahil_toplam_satis, p.para_birimi, rates)
                for p in projects
            ),
            2,
        ),
        try_tahsilat=round(
            sum(_to_try(p.muhasebe.toplam_tahsilat, p.para_birimi, rates) for p in projects), 2
        ),
        try_bakiye=round(
            sum(_to_try(p.muhasebe.kalan_bakiye, p.para_birimi, rates) for p in projects), 2
        ),
        try_net_kar=round(
            sum(_to_try(p.muhasebe.net_kar, p.para_birimi, rates) for p in projects), 2
        ),
        kurlar=rate_rows,
        eksik_kurlar=eksik,
        asama_dagilimi=asama,
        bayi_ozeti=bayi_ozeti,
    )


@router.get("/monthly/export")
async def export_monthly(
    ay: Optional[str] = Query(default=None), user: dict = Depends(require("rapor:disaari"))
):
    report = await monthly_report(ay=ay, user=user)
    rates = await _rates()
    docs = await db.projects.find({"proje_tarihi": {"$regex": f"^{report.ay}"}}).to_list(2000)
    projects = [Project(**_aware(d)) for d in docs]
    labels = await stage_labels()

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F172A")

    def sheet(title: str, headers: list[str], rows: list[list]):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
        for idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(
                14, min(40, len(header) + 6)
            )

    wb.remove(wb.active)

    sheet(
        "Ay Ozeti",
        ["Donem", "Proje Adedi", "Ciro (TRY)", "Tahsilat (TRY)", "Bakiye (TRY)", "Net Kar (TRY)"],
        [
            [
                report.ay,
                report.proje_adet,
                report.try_satis,
                report.try_tahsilat,
                report.try_bakiye,
                report.try_net_kar,
            ]
        ],
    )

    sheet(
        "Kur Bazinda",
        ["Para Birimi", "Proje", "Ciro", "Tahsilat", "Bakiye", "Net Kar", "Kur (TRY)", "Ciro (TRY)"],
        [
            [
                r.para_birimi,
                r.proje_adet,
                r.satis,
                r.tahsilat,
                r.bakiye,
                r.net_kar,
                rates.get(r.para_birimi, 0),
                round(r.satis * rates.get(r.para_birimi, 0), 2),
            ]
            for r in report.kur_dagilimi
        ],
    )

    sheet(
        "Asama Ozeti",
        ["Asama", "Proje Adedi", "Ciro (TRY)"],
        [[a.label, a.adet, a.tutar] for a in report.asama_dagilimi],
    )

    sheet(
        "Bayi Ozeti",
        ["Firma", "Ulke", "Proje Adedi", "Ciro (TRY)", "Tahsilat (TRY)"],
        [[b.firma, b.ulke, b.proje_adet, b.satis_try, b.tahsilat_try] for b in report.bayi_ozeti],
    )

    sheet(
        "Projeler",
        [
            "Proje Kodu",
            "Firma",
            "Musteri",
            "Tarih",
            "Asama",
            "Para Birimi",
            "Ciro",
            "Tahsilat",
            "Bakiye",
            "Ciro (TRY)",
        ],
        [
            [
                p.proje_kodu,
                p.firma,
                p.musteri,
                p.proje_tarihi,
                labels.get(p.durum, p.durum),
                p.para_birimi,
                f"{p.muhasebe.transfer_dahil_toplam_satis:.2f} {p.para_birimi}".strip(),
                f"{p.muhasebe.toplam_tahsilat:.2f} {p.para_birimi}".strip(),
                f"{p.muhasebe.kalan_bakiye:.2f} {p.para_birimi}".strip(),
                round(
                    _to_try(p.muhasebe.transfer_dahil_toplam_satis, p.para_birimi, rates), 2
                ),
            ]
            for p in sorted(projects, key=lambda x: x.proje_tarihi or "")
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="pergola-aylik-{report.ay}.xlsx"'},
    )
